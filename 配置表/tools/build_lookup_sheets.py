# -*- coding: utf-8 -*-
"""
全局参数配置表.xlsx 追加 Monsoon / Trade / War 三个速查 sheet（格式与既有 sheet 完全一致）
- 数值 = 资产当前值（Assets/Game/Resources/MonsoonConfig.asset 等，2026-09-13 快照）
- 嵌套结构展开成行：分组列标明 档位/条目（如 四季·春 / 气候·暴雨 / lootBoxTiers[超高]）
"""
import os
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

BASE = r"D:\unity\Project\Chaotic Chess\配置表"
FILE = os.path.join(BASE, "全局参数配置表.xlsx")

DARK = "1F3864"
LIGHT = "F2F5FA"
BLUE = "000000FF"
WHITE = "FFFFFF"
HEADERS = ["分组", "参数名", "类型", "默认值", "说明"]
WIDTHS = [14, 34, 10, 24, 52]


def make_sheet(wb, title, rows):
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = title
    c.fill = PatternFill("solid", fgColor=DARK)
    c.font = Font(bold=True, color=WHITE, size=12)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24
    for ci, h in enumerate(HEADERS, start=1):
        cc = ws.cell(row=2, column=ci, value=h)
        cc.fill = PatternFill("solid", fgColor=DARK)
        cc.font = Font(bold=True, color=WHITE, size=11)
        cc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    for ri, row in enumerate(rows, start=3):
        fill = PatternFill("solid", fgColor=LIGHT) if ri % 2 == 0 else None
        for ci, v in enumerate(row, start=1):
            cc = ws.cell(row=ri, column=ci, value=v)
            if fill:
                cc.fill = fill
            if ci == 1:
                cc.font = Font(bold=True, color=DARK, size=10)
            elif ci == 2:
                cc.font = Font(color=BLUE, size=10)
            else:
                cc.font = Font(size=10)
            cc.alignment = Alignment(vertical="top", wrap_text=True)
    for ci, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[chr(64 + ci)].width = w
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:E{len(rows) + 2}"
    return ws


# ============ Monsoon ============
R = "（引用型：Inspector 配置）"
monsoon_rows = [
    # 四季（seasons[]）
    ("四季·春", "slotName", "string", "春", "季节槽位名（seasons[0]）"),
    ("四季·春", "effectType", "EffectType", "Heal", "效果类型：回血"),
    ("四季·春", "value", "int", "1", "该方回合开始全体存活棋子回 1 血（封顶 MaxHP）"),
    ("四季·夏", "slotName", "string", "夏", "季节槽位名（seasons[1]）"),
    ("四季·夏", "effectType", "EffectType", "Attack", "效果类型：攻击"),
    ("四季·夏", "value", "int", "2", "有效攻击力聚合 +2（双方全体）"),
    ("四季·秋", "slotName", "string", "秋", "季节槽位名（seasons[2]）"),
    ("四季·秋", "effectType", "EffectType", "Gold", "效果类型：金币"),
    ("四季·秋", "value", "int", "10", "该方回合结束 +10 金币"),
    ("四季·冬", "slotName", "string", "冬", "季节槽位名（seasons[3]）"),
    ("四季·冬", "effectType", "EffectType", "Damage", "效果类型：扣血"),
    ("四季·冬", "value", "int", "1", "该方回合结束全体存活棋子 -1 血（可致死）"),
    # 昼夜
    ("昼夜·昼", "slotName", "string", "昼", "昼槽位名"),
    ("昼夜·昼", "effectType", "EffectType", "Move", "效果类型：移动"),
    ("昼夜·昼", "value", "int", "1", "移动范围聚合 +1"),
    ("昼夜·夜", "slotName", "string", "夜", "夜槽位名"),
    ("昼夜·夜", "effectType", "EffectType", "Move", "效果类型：移动"),
    ("昼夜·夜", "value", "int", "-1", "移动范围聚合 -1（clamp 到 minMoveRange）"),
    ("昼夜", "minMoveRange", "int", "1", "移动范围下限（夜 -1 后至少可行动 1 格）"),
    # 节奏
    ("节奏", "roundsPerSeason", "int", "6", "每多少轮切换一季"),
    ("节奏", "dayRoundsPerSeason", "int", "3", "每季前多少轮为昼（其余为夜）"),
    # 气候条目（climates[]）
    ("气候·暴雨", "climateName", "string", "暴雨", "气候条目名（climates[0]）"),
    ("气候·暴雨", "probability", "float", "0.05", "每轮开始出现概率（0~1）"),
    ("气候·暴雨", "restrictedSeasonIndex", "int", "-1", "限制季节下标（-1=不限；0春 1夏 2秋 3冬）"),
    ("气候·暴雨", "effectKind", "ClimateEffectKind", "Rain", "效果类型：暴雨"),
    ("气候·暴雨", "exclusive", "bool", "false", "是否独占（高温=true 时本轮不再判定其他）"),
    ("气候·暴雪", "climateName", "string", "暴雪", "气候条目名（climates[1]）"),
    ("气候·暴雪", "probability", "float", "0.05", "每轮开始出现概率（0~1）"),
    ("气候·暴雪", "restrictedSeasonIndex", "int", "3", "限制季节下标：冬季"),
    ("气候·暴雪", "effectKind", "ClimateEffectKind", "Blizzard", "效果类型：暴雪"),
    ("气候·暴雪", "exclusive", "bool", "false", "是否独占"),
    ("气候·高温", "climateName", "string", "高温", "气候条目名（climates[2]）"),
    ("气候·高温", "probability", "float", "0.05", "每轮开始出现概率（0~1）"),
    ("气候·高温", "restrictedSeasonIndex", "int", "1", "限制季节下标：夏季"),
    ("气候·高温", "effectKind", "ClimateEffectKind", "Heat", "效果类型：高温（行动一次后锁定，无数值参数）"),
    ("气候·高温", "exclusive", "bool", "true", "是否独占"),
    ("气候·飓风", "climateName", "string", "飓风", "气候条目名（climates[3]）"),
    ("气候·飓风", "probability", "float", "0.163", "每轮开始出现概率（0~1）"),
    ("气候·飓风", "restrictedSeasonIndex", "int", "-1", "限制季节下标（-1=不限）"),
    ("气候·飓风", "effectKind", "ClimateEffectKind", "Hurricane", "效果类型：飓风（商店关闭+掉幸运方块）"),
    ("气候·飓风", "exclusive", "bool", "false", "是否独占"),
    ("气候·雷暴", "climateName", "string", "雷暴", "气候条目名（climates[4]）"),
    ("气候·雷暴", "probability", "float", "0.168", "每轮开始出现概率（0~1）"),
    ("气候·雷暴", "restrictedSeasonIndex", "int", "-1", "限制季节下标（-1=不限）"),
    ("气候·雷暴", "effectKind", "ClimateEffectKind", "Lightning", "效果类型：雷暴（雷劈+雷电残留）"),
    ("气候·雷暴", "exclusive", "bool", "false", "是否独占"),
    # 气候通用
    ("气候", "climateDurationMinRounds", "int", "1", "气候持续轮数下限"),
    ("气候", "climateDurationMaxRounds", "int", "3", "气候持续轮数上限"),
    # 暴雨/暴雪参数
    ("暴雨参数", "rainAttackRangePenalty", "int", "1", "暴雨：上回合受击棋子本回合攻击距离减量（下限 1 格）"),
    ("暴雪参数", "blizzardKnockbackDistance", "int", "1", "暴雪：被攻击棋子被击退格数"),
    # 飓风参数
    ("飓风参数", "hurricaneDropChancePerRound", "float", "0.5", "每轮掉方块概率（0~1）"),
    ("飓风参数", "hurricaneDropCountMin", "int", "3", "每批掉落数量下限"),
    ("飓风参数", "hurricaneDropCountMax", "int", "4", "每批掉落数量上限"),
    ("飓风参数", "lootBoxPrefab", "GameObject（引用）", R, "幸运方块 3D 实体 prefab"),
    ("飓风参数", "lootBoxYOffset", "float", "0.38", "方块 pivot 到脚底的 Y 距离"),
    # lootBoxTiers
    ("lootBoxTiers[0]·超高", "tierName", "string", "超高", "档位名（日志用）"),
    ("lootBoxTiers[0]·超高", "material", "Material（引用）", R, "档位方块材质（品质视觉）"),
    ("lootBoxTiers[0]·超高", "weight", "float", "0.05", "档位概率权重"),
    ("lootBoxTiers[0]·超高", "includeNormalEquipment", "bool", "true", "是否含普通装备来源"),
    ("lootBoxTiers[0]·超高", "equipmentTier", "EquipmentTier", "Advanced", "普通装备 tier"),
    ("lootBoxTiers[0]·超高", "minCityPrice", "int", "50", "城邦装备/道具价格下限"),
    ("lootBoxTiers[0]·超高", "maxCityPrice", "int", "9999", "城邦装备/道具价格上限"),
    ("lootBoxTiers[0]·超高", "goldMin", "int", "100", "金币下限"),
    ("lootBoxTiers[0]·超高", "goldMax", "int", "150", "金币上限"),
    ("lootBoxTiers[1]·高", "tierName", "string", "高", "档位名（日志用）"),
    ("lootBoxTiers[1]·高", "material", "Material（引用）", R, "档位方块材质"),
    ("lootBoxTiers[1]·高", "weight", "float", "0.15", "档位概率权重"),
    ("lootBoxTiers[1]·高", "includeNormalEquipment", "bool", "true", "是否含普通装备来源"),
    ("lootBoxTiers[1]·高", "equipmentTier", "EquipmentTier", "Intermediate", "普通装备 tier"),
    ("lootBoxTiers[1]·高", "minCityPrice", "int", "50", "城邦装备/道具价格下限"),
    ("lootBoxTiers[1]·高", "maxCityPrice", "int", "9999", "城邦装备/道具价格上限"),
    ("lootBoxTiers[1]·高", "goldMin", "int", "20", "金币下限"),
    ("lootBoxTiers[1]·高", "goldMax", "int", "60", "金币上限"),
    ("lootBoxTiers[2]·中", "tierName", "string", "中", "档位名（日志用）"),
    ("lootBoxTiers[2]·中", "material", "Material（引用）", R, "档位方块材质"),
    ("lootBoxTiers[2]·中", "weight", "float", "0.3", "档位概率权重"),
    ("lootBoxTiers[2]·中", "includeNormalEquipment", "bool", "true", "是否含普通装备来源"),
    ("lootBoxTiers[2]·中", "equipmentTier", "EquipmentTier", "Basic", "普通装备 tier"),
    ("lootBoxTiers[2]·中", "minCityPrice", "int", "0", "城邦装备/道具价格下限"),
    ("lootBoxTiers[2]·中", "maxCityPrice", "int", "49", "城邦装备/道具价格上限"),
    ("lootBoxTiers[2]·中", "goldMin", "int", "10", "金币下限"),
    ("lootBoxTiers[2]·中", "goldMax", "int", "20", "金币上限"),
    ("lootBoxTiers[3]·低", "tierName", "string", "低", "档位名（日志用）"),
    ("lootBoxTiers[3]·低", "material", "Material（引用）", R, "档位方块材质"),
    ("lootBoxTiers[3]·低", "weight", "float", "0.5", "档位概率权重"),
    ("lootBoxTiers[3]·低", "includeNormalEquipment", "bool", "false", "是否含普通装备来源（低档纯金币）"),
    ("lootBoxTiers[3]·低", "equipmentTier", "EquipmentTier", "Basic", "普通装备 tier（含普通装备=false 时忽略）"),
    ("lootBoxTiers[3]·低", "minCityPrice", "int", "0", "城邦装备/道具价格下限"),
    ("lootBoxTiers[3]·低", "maxCityPrice", "int", "-1", "城邦装备/道具价格上限（max<min=来源为空）"),
    ("lootBoxTiers[3]·低", "goldMin", "int", "1", "金币下限"),
    ("lootBoxTiers[3]·低", "goldMax", "int", "6", "金币上限"),
    # 雷暴参数
    ("雷暴参数", "lightningStrikeMinCount", "int", "2", "雷劈格子数量下限"),
    ("雷暴参数", "lightningStrikeMaxCount", "int", "4", "雷劈格子数量上限"),
    ("雷暴参数", "lightningStrikeDamage", "int", "20", "劈中棋子真实伤害（无视防御）"),
    ("雷暴参数", "lightningResidueInitialPower", "int", "8", "雷电残留初始强度"),
    ("雷暴参数", "lightningResidueMaterial", "Material（引用）", R, "雷电残留格 3D 材质"),
    ("雷暴参数", "lightningResidueDecayPerRound", "int", "2", "残留强度每轮衰减量"),
    ("雷暴参数", "lightningPermanentBonusChance", "float", "1", "雷劈强化概率（0~1）"),
    ("雷暴参数", "lightningResiduePermanentBonusChance", "float", "0.6", "残留强化概率（0~1）"),
    # 雷劈强化池
    ("雷暴·雷劈强化池[0]", "type", "PermanentBonusType", "HP", "属性类型：MaxHP +N（当前血量不同步涨）"),
    ("雷暴·雷劈强化池[0]", "value", "int", "10", "加成数值（可叠加）"),
    ("雷暴·雷劈强化池[0]", "weight", "float", "4", "抽中概率权重（非归一化）"),
    ("雷暴·雷劈强化池[1]", "type", "PermanentBonusType", "Attack", "属性类型：EffectiveAttack +N"),
    ("雷暴·雷劈强化池[1]", "value", "int", "10", "加成数值"),
    ("雷暴·雷劈强化池[1]", "weight", "float", "3", "抽中概率权重"),
    ("雷暴·雷劈强化池[2]", "type", "PermanentBonusType", "Defense", "属性类型：EffectiveDefense +N"),
    ("雷暴·雷劈强化池[2]", "value", "int", "5", "加成数值"),
    ("雷暴·雷劈强化池[2]", "weight", "float", "3", "抽中概率权重"),
    ("雷暴·雷劈强化池[3]", "type", "PermanentBonusType", "Move", "属性类型：MoveRange +N"),
    ("雷暴·雷劈强化池[3]", "value", "int", "1", "加成数值"),
    ("雷暴·雷劈强化池[3]", "weight", "float", "1", "抽中概率权重"),
    ("雷暴·雷劈强化池[4]", "type", "PermanentBonusType", "Range", "属性类型：AttackRange +N"),
    ("雷暴·雷劈强化池[4]", "value", "int", "1", "加成数值"),
    ("雷暴·雷劈强化池[4]", "weight", "float", "1", "抽中概率权重"),
    # 残留强化池
    ("雷暴·残留强化池[0]", "type", "PermanentBonusType", "HP", "属性类型：MaxHP +N"),
    ("雷暴·残留强化池[0]", "value", "int", "4", "加成数值"),
    ("雷暴·残留强化池[0]", "weight", "float", "6", "抽中概率权重"),
    ("雷暴·残留强化池[1]", "type", "PermanentBonusType", "Attack", "属性类型：EffectiveAttack +N"),
    ("雷暴·残留强化池[1]", "value", "int", "4", "加成数值"),
    ("雷暴·残留强化池[1]", "weight", "float", "4", "抽中概率权重"),
    ("雷暴·残留强化池[2]", "type", "PermanentBonusType", "Defense", "属性类型：EffectiveDefense +N"),
    ("雷暴·残留强化池[2]", "value", "int", "2", "加成数值"),
    ("雷暴·残留强化池[2]", "weight", "float", "4", "抽中概率权重"),
    ("雷暴·残留强化池[3]", "type", "PermanentBonusType", "Move", "属性类型：MoveRange +N"),
    ("雷暴·残留强化池[3]", "value", "int", "1", "加成数值"),
    ("雷暴·残留强化池[3]", "weight", "float", "1", "抽中概率权重"),
    ("雷暴·残留强化池[4]", "type", "PermanentBonusType", "Range", "属性类型：AttackRange +N"),
    ("雷暴·残留强化池[4]", "value", "int", "1", "加成数值"),
    ("雷暴·残留强化池[4]", "weight", "float", "1", "抽中概率权重"),
]

# ============ Trade ============
trade_rows = [
    ("透支", "overdraftFloor", "int", "-50", "金币透支下限（买装备/买道具/拍卖成交可透支到该下限）"),
    ("拍卖", "auctionDealCountdownTurns", "int", "6", "拍卖成交倒计时（回合）：出价后每回合结束 -1，归零成交"),
    ("拍卖", "auctionRefreshIntervalTurns", "int", "5", "拍卖行刷新间隔（回合）：每 N 回合从物品池补 1 件新拍品"),
    ("利息与信誉", "interestGoldInterval", "int", "10", "利息结算间隔（金币）：每轮结束每拥有 N 金币得 1 金币利息"),
    ("利息与信誉", "debtCreditTickTurns", "int", "8", "欠钱扣信誉计次（回合）：持续欠钱满 N 回合扣一次信誉"),
    ("利息与信誉", "auctionReputationThreshold", "int", "6", "拍卖信誉门槛：信誉低于该值无法在拍卖行出价"),
    ("利息与信誉", "shopPricePercentPerReputation", "int", "10", "商店涨价：每失去 1 点信誉价格上涨百分比（信誉 10=原价）"),
    ("地下交易", "undergroundDamageThreshold", "int", "60", "地下交易激活阈值（累计直接伤害）"),
    ("地下交易", "undergroundProbBasePercent", "int", "10", "地下交易触发基础概率（%）"),
    ("地下交易", "undergroundProbStepDamage", "int", "30", "概率步进伤害：累计伤害每超阈值该值概率 +stepPercent"),
    ("地下交易", "undergroundProbStepPercent", "int", "20", "概率步进（%）：每步增加概率（封顶 100%）"),
    ("地下交易", "undergroundReputationThreshold", "int", "6", "地下交易信誉门槛：信誉低于该值才可激活/维持"),
]

# ============ War ============
war_rows = [
    ("主将加成", "generalHPBonus", "int", "20", "主将生命上限加成（登记时同步补满等量当前血量）"),
    ("主将加成", "generalDefenseBonus", "int", "3", "主将防御加成"),
    ("主将加成", "generalAttackBonus", "int", "6", "主将攻击加成"),
    ("主将加成", "generalMoveBonus", "int", "1", "主将移动距离加成（不含攻击距离）"),
    ("主将装备额外加成", "generalEquipStatBonus", "int", "2", "主将装备每实际提供一项属性，该项额外加该值"),
    ("棋盘扩展", "boardExpansionRings", "int", "2", "战争之城时棋盘向外扩展的环数（0=不扩展）"),
]


def main():
    wb = openpyxl.load_workbook(FILE)
    make_sheet(wb, "Monsoon", monsoon_rows)
    make_sheet(wb, "Trade", trade_rows)
    make_sheet(wb, "War", war_rows)
    wb.save(FILE)
    print(f"[lookup] 已追加 sheet：Monsoon {len(monsoon_rows)} 行 / Trade {len(trade_rows)} 行 / War {len(war_rows)} 行")
    print(f"[lookup] sheet 列表：{wb.sheetnames}")


if __name__ == "__main__":
    main()
