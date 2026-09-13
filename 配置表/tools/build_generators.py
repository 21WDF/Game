# -*- coding: utf-8 -*-
"""
装备 / 道具生成器 xlsx 构建脚本（一次性构建，可重复运行刷新现值）
- 读 Assets/Game/Resources/Equipment 与 GridItem 下所有 EquipmentData / GridItemData 资产
- 按「全局参数配置表」格式规范生成：首行标题(深底白字合并) / 第二行表头 / 数据行隔行底色 /
  输入列蓝字 / 冻结前两行 / 自动筛选
- 预填当前资产值（表格不变时回写脚本不产生任何改动 → 幂等）
"""
import codecs
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ASSETS = r"D:\unity\Project\Chaotic Chess\Assets"
OUT = r"D:\unity\Project\Chaotic Chess\配置表"

DARK = "1F3864"      # 标题/表头深色底
LIGHT = "F2F5FA"     # 数据行隔行底色
BLUE = "000000FF"    # 输入列蓝字
GREY = "595959"
WHITE = "FFFFFF"

EQUIP_TIER = {0: "Basic", 1: "Intermediate", 2: "Advanced", 3: "CityState"}
CITY_KIND = {0: "None", 1: "Trade", 2: "Merriment", 3: "Monsoon", 4: "Occult", 5: "War", 6: "Element"}


def decode_str(raw):
    """解码 Unity YAML 字符串值（去掉引号，解码 \\uXXXX）"""
    s = raw.strip()
    if len(s) >= 2 and s[0] == '"' and s[-1] == '"':
        return codecs.decode(s[1:-1], "unicode_escape")
    if len(s) >= 2 and s[0] == "'" and s[-1] == "'":
        return codecs.decode(s[1:-1], "unicode_escape")
    return s


def parse_asset(path):
    """解析单个 .asset：返回 (class_id, fields) ；fields 顶层标量 + passives 列表"""
    with open(path, encoding="utf-8") as f:
        text = f.read()
    lines = text.split("\n")
    class_id = ""
    j = 0
    while j < len(lines) and "MonoBehaviour:" not in lines[j]:
        j += 1
    if j >= len(lines):
        return class_id, {}
    j += 1
    fields = {}
    passives = []
    cur = None
    in_passives = False
    for line in lines[j:]:
        if line.startswith("---"):
            break
        if not line.strip():
            continue
        stripped = line.strip()
        indent = len(line) - len(line.lstrip())
        if indent == 2 and stripped.startswith("- "):
            in_passives = True
            cur = {}
            passives.append(cur)
            # 兼容 `- className: X` 内联首字段（dash 与字段同行）
            rest = stripped[2:].strip()
            m2 = re.match(r"^([A-Za-z_]\w*):\s*(.*)$", rest)
            if m2:
                cur[m2.group(1)] = m2.group(2)
            continue
        if indent == 2:
            m = re.match(r"^([A-Za-z_]\w*):\s*(.*)$", stripped)
            if m:
                key, val = m.group(1), m.group(2)
                if key == "m_EditorClassIdentifier":
                    class_id = val.strip()
                elif key == "passives":
                    if val.strip() == "[]":
                        passives = []
                    in_passives = True
                    cur = None
                else:
                    fields[key] = val
                    in_passives = False
            continue
        if indent >= 4 and in_passives and cur is not None:
            m = re.match(r"^([A-Za-z_]\w*):\s*(.*)$", stripped)
            if m:
                cur[m.group(1)] = m.group(2)
    if passives:
        fields["passives"] = passives
    return class_id, fields


def enum_name(mapping, raw):
    """enum 字段：原始行值可能缺失(默认0) 或为数字"""
    if raw is None or raw.strip() == "":
        return mapping[0]
    try:
        return mapping[int(raw)]
    except ValueError:
        return raw.strip()


def int_val(raw, default=0):
    if raw is None or raw.strip() == "":
        return default
    try:
        return int(raw)
    except ValueError:
        return default


def bool_txt(raw):
    if raw is None or raw.strip() == "":
        return "False"
    try:
        return "True" if int(raw) else "False"
    except ValueError:
        return str(raw).strip()


def fmt_passives(passives):
    """passives → ';' 分隔多个被动，子字段 '|' 分隔 className|jsonParams|passiveName|isUnique|uniqueId"""
    if not passives:
        return ""
    parts = []
    for p in passives:
        name = decode_str(p.get("className", "")) if "className" in p else ""
        jp = decode_str(p.get("jsonParams", "")) if "jsonParams" in p else ""
        pn = decode_str(p.get("passiveName", "")) if "passiveName" in p else ""
        iu = "True" if int_val(p.get("isUnique", "0")) else "False"
        uid = decode_str(p.get("uniqueId", "")) if "uniqueId" in p else ""
        parts.append("|".join([name, jp, pn, iu, uid]))
    return ";".join(parts)


# ---------- 样式 ----------
def style_workbook(wb, title, headers, rows, widths, blue_cols, ref_cols):
    ws = wb.active
    last_col = len(headers)
    last_letter = get_column_letter(last_col)
    # 首行标题（合并 + 深色底白字）
    ws.merge_cells(f"A1:{last_letter}1")
    c = ws["A1"]
    c.value = title
    c.fill = PatternFill("solid", fgColor=DARK)
    c.font = Font(bold=True, color=WHITE, size=12)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24
    # 第二行表头
    for ci, h in enumerate(headers, start=1):
        cc = ws.cell(row=2, column=ci, value=h)
        cc.fill = PatternFill("solid", fgColor=DARK)
        cc.font = Font(bold=True, color=WHITE, size=11)
        cc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30
    # 数据行（隔行底色 + 输入列蓝字）
    for ri, row in enumerate(rows, start=3):
        fill = PatternFill("solid", fgColor=LIGHT) if ri % 2 == 0 else None
        for ci, v in enumerate(row, start=1):
            cc = ws.cell(row=ri, column=ci, value=v if v is not None else "")
            if fill:
                cc.fill = fill
            if ci in ref_cols:
                cc.font = Font(color=GREY, italic=True, size=9)
            elif ci in blue_cols:
                cc.font = Font(color=BLUE, size=10)
            else:
                cc.font = Font(size=10)
            cc.alignment = Alignment(vertical="top", wrap_text=True)
    # 列宽 / 冻结 / 筛选
    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{last_letter}{len(rows) + 2}"
    return ws


# ================= 装备生成器 =================
def build_equipment():
    items = []
    base = os.path.join(ASSETS, "Game", "Resources", "Equipment")
    for root, _dirs, files in os.walk(base):
        for fn in sorted(files):
            if not fn.endswith(".asset"):
                continue
            full = os.path.join(root, fn)
            rel = os.path.relpath(full, ASSETS).replace("\\", "/")
            class_id, f = parse_asset(full)
            if "EquipmentData" not in class_id:
                continue  # 跳过 AuctionItemData / UndergroundItemData 包装资产
            tier = enum_name(EQUIP_TIER, f.get("tier"))
            city = enum_name(CITY_KIND, f.get("cityState"))
            items.append({
                "path": rel,
                "id": int_val(f.get("id")),
                "name": decode_str(f.get("displayName", f.get("m_Name", ""))),
                "desc": decode_str(f.get("description", "")),
                "tier": tier,
                "city": city,
                "price": int_val(f.get("price")),
                "bonusHP": int_val(f.get("bonusHP")),
                "bonusAttack": int_val(f.get("bonusAttack")),
                "bonusDefense": int_val(f.get("bonusDefense")),
                "bonusMove": int_val(f.get("bonusMoveRange")),
                "bonusRange": int_val(f.get("bonusAttackRange")),
                "bonusEnergy": int_val(f.get("bonusEnergyPerTurn")),
                "bonusAPCap": int_val(f.get("bonusAPCap")),
                "passives": fmt_passives(f.get("passives")),
            })
    items.sort(key=lambda x: x["id"])
    headers = ["目标资产路径（相对 Assets）", "id", "displayName", "description", "tier", "cityState",
               "price", "icon（引用型，请在 Inspector 修改）", "bonusHP", "bonusAttack", "bonusDefense",
               "bonusMoveRange", "bonusAttackRange", "bonusEnergyPerTurn", "bonusAPCap",
               "passives（; 分隔多个被动，子字段按 className|jsonParams|passiveName|isUnique|uniqueId 用 | 分隔，可留空段）", "说明"]
    rows = []
    for it in items:
        rows.append([
            it["path"], it["id"], it["name"], it["desc"], it["tier"], it["city"], it["price"],
            "（引用型：请在 Inspector 修改）",
            it["bonusHP"], it["bonusAttack"], it["bonusDefense"], it["bonusMove"], it["bonusRange"],
            it["bonusEnergy"], it["bonusAPCap"], it["passives"], "",
        ])
    widths = [46, 6, 14, 22, 14, 12, 8, 22, 8, 8, 8, 8, 8, 8, 8, 48, 34]
    wb = Workbook()
    blue = set(range(2, 17)) | {17}
    blue.discard(8)  # icon 引用列非输入列
    style_workbook(wb, f"装备配置生成器 · 一行一件装备 · 共 {len(rows)} 件 · 资产目录：Assets/Game/Resources/Equipment",
                   headers, rows, widths, blue, ref_cols={8, 17})
    wb.save(os.path.join(OUT, "装备配置生成器.xlsx"))
    print(f"[build] 装备配置生成器.xlsx：{len(rows)} 行")


# ================= 道具生成器 =================
def build_griditem():
    items = []
    base = os.path.join(ASSETS, "Game", "Resources", "GridItem")
    for fn in sorted(os.listdir(base)):
        if not fn.endswith(".asset"):
            continue
        full = os.path.join(base, fn)
        rel = os.path.relpath(full, ASSETS).replace("\\", "/")
        class_id, f = parse_asset(full)
        if "GridItemData" not in class_id:
            continue
        items.append({
            "path": rel,
            "id": int_val(f.get("id")),
            "name": decode_str(f.get("displayName", f.get("m_Name", ""))),
            "price": int_val(f.get("price")),
            "apCost": int_val(f.get("apCost"), 1),
            "effectClass": decode_str(f.get("effectClassName", "")),
            "effectJson": decode_str(f.get("effectJsonParams", "")),
            "reqPiece": bool_txt(f.get("requiresSelectedPiece")),
            "minTiles": int_val(f.get("minTilesAfter"), 7),
        })
    items.sort(key=lambda x: x["id"])
    headers = ["目标资产路径（相对 Assets）", "id", "displayName", "icon（引用型，请在 Inspector 修改）",
               "price", "apCost", "effectClassName", "effectJsonParams", "requiresSelectedPiece",
               "minTilesAfter", "说明"]
    rows = []
    for it in items:
        rows.append([
            it["path"], it["id"], it["name"], "（引用型：请在 Inspector 修改）",
            it["price"], it["apCost"], it["effectClass"], it["effectJson"], it["reqPiece"],
            it["minTiles"], "",
        ])
    widths = [46, 6, 14, 22, 8, 8, 22, 28, 14, 12, 34]
    wb = Workbook()
    blue = set(range(2, 11)) | {11}
    blue.discard(4)
    style_workbook(wb, f"道具配置生成器 · 一行一件道具 · 共 {len(rows)} 件 · 资产目录：Assets/Game/Resources/GridItem",
                   headers, rows, widths, blue, ref_cols={4, 11})
    wb.save(os.path.join(OUT, "道具配置生成器.xlsx"))
    print(f"[build] 道具配置生成器.xlsx：{len(rows)} 行")


if __name__ == "__main__":
    build_equipment()
    build_griditem()
    print("[build] 完成")
