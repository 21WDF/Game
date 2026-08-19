# -*- coding: utf-8 -*-
"""生成「棋子配置生成器.xlsx」—— 九棋子（杜林/菲林斯/克洛琳德/哥伦比亚/甘雨/安柏/莫娜/爱可菲/雷电将军）+ 兔兔伯爵傀儡 统一配置表。
按类别分组：被动配置生成器 / 大招配置生成器 / 移动配置生成器 / 攻击配置生成器 / 效果速查 / 配置清单。
复用 gen_dulin_config.py 的样式与公式写法（CHAR(34) 生成 JSON 引号）。
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = r"D:\unity\Project\Chaotic Chess\棋子配置生成器.xlsx"

# ---------- 样式 ----------
F_TITLE = Font(name="微软雅黑", size=16, bold=True, color="1F3864")
F_SUB   = Font(name="微软雅黑", size=10, color="595959")
F_BLOCK = Font(name="微软雅黑", size=12, bold=True, color="FFFFFF")
F_HEAD  = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
F_BODY  = Font(name="微软雅黑", size=10, color="222222")
F_FIELD = Font(name="Consolas", size=10, bold=True, color="1F3864")
F_GEN   = Font(name="Consolas", size=10, bold=True, color="006100")

FILL_BLOCK = PatternFill("solid", fgColor="4472C4")
FILL_HEAD  = PatternFill("solid", fgColor="5B9BD5")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_DEF   = PatternFill("solid", fgColor="F2F2F2")
FILL_GEN   = PatternFill("solid", fgColor="E2EFDA")
FILL_TITLE = PatternFill("solid", fgColor="D9E1F2")

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left", vertical="center", wrap_text=True)


# ---------- 被动数据 ----------
# (棋子名, 效果名, className, 说明, [(字段, 类型, 标签, 默认, 注释)])
PASSIVES = [
    ("杜林（火）", "炎星溅射", "FlameNovaPassive",
     "普攻命中后对自身周围 splashRadius 格内敌方造成物理伤害（带火元素，触发反应）",
     [("baseDamage", "int", "基础伤害", 0, "溅射基础伤害数值"),
      ("attackPercent", "int", "攻击力百分比", 50, "按攻击力比例溅射(0-100)"),
      ("splashRadius", "int", "溅射半径", 2, "以自身为圆心的溅射半径(格)"),
      ("goldPercent", "float", "金币比例", 0, "溅射伤害转金币比例(0=不产金币)"),
      ("elemental", "bool", "带元素", "true", "是否附着火元素触发反应")]),
    ("菲林斯（雷）", "雷蚀强化", "SparkEmpowerPassive",
     "普攻命中后进入强化状态持续 duration 回合，强化态改变大招形态",
     [("duration", "int", "持续回合", 6, "强化状态持续回合数")]),
    ("克洛琳德（雷）", "雷径惩戒", "VoltPathPassive",
     "移动路径飞越的敌方棋子受到物理伤害（带雷元素，触发超导/感电）",
     [("baseDamage", "int", "基础伤害", 0, "路径伤害基础数值"),
      ("attackPercent", "int", "攻击力百分比", 0, "按攻击力比例伤害(0-100)"),
      ("elemental", "bool", "带元素", "true", "是否附着雷元素触发反应")]),
    ("哥伦比亚（水）", "生命之息", "RevivePassive",
     "己方棋子死亡时立即原地复活 reviveHp 血（一局仅触发一次）",
     [("reviveHp", "int", "复活血量", 10, "复活后的生命值")]),
    ("安柏（火）", "快速射击", "DoubleStrikePassive",
     "普攻造成两次物理伤害（第二段 percent% 攻击力，完整管道、不重复充能）",
     [("percent", "int", "二段百分比", 100, "第二段攻击力百分比(100=完整一段)")]),
    ("兔兔伯爵（傀儡）", "傀儡核心", "PuppetPassive",
     "傀儡行为核心：受击计数引爆 / 持续回合引爆 / 嘲讽光环 / 爆炸（物理+主人火元素）",
     [("hitsToDetonate", "int", "受击引爆阈值", 3, "受击次数满则引爆"),
      ("maxTurns", "int", "持续回合", 6, "持续回合满则引爆"),
      ("tauntRange", "int", "嘲讽半径", 2, "距傀儡≤此距离的敌方只能攻击傀儡"),
      ("explosionRadius", "int", "爆炸半径", 1, "引爆对周围格伤害范围"),
      ("baseDamage", "int", "爆炸基础伤害", 0, "引爆基础伤害数值"),
      ("attackPercent", "int", "攻击力百分比", 100, "按主人攻击力比例(0-100)")]),
    ("莫娜（水）", "水域领域", "DomainPassive",
     "普攻命中后在目标位置留下半径 radius 格领域，每回合对领域内敌方造成物理伤害（带水元素）；再普攻移动领域",
     [("radius", "int", "领域半径", 2, "领域作用半径(格)"),
      ("baseDamage", "int", "基础伤害", 0, "每回合基础伤害数值"),
      ("attackPercent", "int", "攻击力百分比", 50, "按攻击力比例(0-100)")]),
    ("爱可菲（冰）", "冰锋蚀甲", "FrostBreakPassive",
     "受到爱可菲攻击的敌方防御值降低 amount（不可叠加=刷新），持续 duration 回合",
     [("amount", "int", "减防值", 4, "防御降低固定值(不可叠加)"),
      ("duration", "int", "持续回合", 4, "减防持续回合数")]),
    ("爱可菲（冰）", "凛冬风暴", "BlizzardPassive",
     "由大招激活的持续段：每回合对周围 radius 格内「最近+最远」各1个敌方造成魔法伤害（带冰元素）",
     [("radius", "int", "风暴半径", 3, "以爱可菲当前位置为圆心(格)"),
      ("baseDamage", "int", "基础伤害", 0, "每回合魔法基础伤害"),
      ("attackPercent", "int", "攻击力百分比", 30, "按攻击力比例(0-100)")]),
    ("雷电将军（雷）", "雷罚恶曜之眼", "StormEyePassive",
     "全局回能：己方其他棋子释放大招时，雷电将军回复 energyGain 点能量（不含自己）",
     [("energyGain", "int", "回能量", 10, "队友开大时回能量值")]),
    ("雷电将军（雷）", "协同状态", "CoordinatedStrikePassive",
     "由大招激活：雷电将军普攻全体己方回能；己方攻击命中触发远程协同攻击（含自己、防递归）",
     [("baseDamage", "int", "协同基础伤害", 0, "协同攻击基础伤害数值"),
      ("attackPercent", "int", "攻击力百分比", 50, "按雷电将军攻击力比例(0-100)"),
      ("energyPerAttack", "int", "普攻回能量", 5, "雷电将军普攻时全体己方回能量")]),
]

# ---------- 大招数据 ----------
# (棋子名, 效果名, className, 说明, [(字段, 类型, 标签, 默认, 注释)])
ULTIMATES = [
    ("杜林（火）", "贯日炎枪", "FlameLanceUltimate",
     "指定格模式(Tile)，点格子定方向，沿最近 hex 方向直线穿透 length 格内敌方（含火元素反应）",
     [("length", "int", "射线长度", 6, "直线穿透格数(需与 tileTargetRange 一致)"),
      ("bonusDamage", "int", "附加攻击力", 0, "附加在 EffectiveAttack 上的额外攻击力")]),
    ("菲林斯（雷）", "雷殛横光", "ThunderSweepUltimate",
     "指定格模式(Tile)：普通态横向3格物理；强化态周围1格魔法（强化态消耗 empoweredEnergyCost 能量、普通态全清）",
     [("normalBonusDamage", "int", "普通态附加攻击", 0, "普通态附加在 EffectiveAttack 上的攻击力"),
      ("empoweredBaseDamage", "int", "强化态魔法基础", 0, "强化态魔法基础伤害值"),
      ("empoweredAttackPercent", "int", "强化态攻击百分比", 0, "强化态按攻击力比例(0-100)"),
      ("empoweredEnergyCost", "int", "强化态能量消耗", 50, "强化态消耗的能量量(保留剩余)"),
      ("empoweredElemental", "bool", "强化态带元素", "true", "强化态魔法段是否附着雷元素")]),
    ("克洛琳德（雷）", "雷驰突进", "ThunderDashUltimate",
     "指定格模式(Tile+Ray)，向指定方向直线冲刺并回复 healAmount 生命（复用移动管道，触发路径伤害）",
     [("healAmount", "int", "回血量", 10, "冲刺后回复的生命值")]),
    ("哥伦比亚（水）", "潮涌战意", "TideSurgeUltimate",
     "自身增益(Self)，周围 radius 格内己方攻击力 +bonusAttack 持续 duration 回合",
     [("radius", "int", "增益半径", 5, "以自身为圆心的增益半径(格)"),
      ("bonusAttack", "int", "攻击加成", 0, "攻击力加成值"),
      ("duration", "int", "持续回合", 4, "buff 持续回合数")]),
    ("甘雨（冰）", "霜华贯矢", "FrostArrowUltimate",
     "指定格模式(Tile)，按蓄力层数三档：1层单点物理 / 2层周围1格物理 / 3层周围2格魔法（不耗能量，消耗蓄力层数）",
     [("physicalBonusDamage", "int", "物理附加攻击", 0, "1/2层物理段附加攻击力"),
      ("magicBaseDamage", "int", "魔法基础伤害", 0, "3层魔法基础伤害值"),
      ("magicAttackPercent", "int", "魔法攻击百分比", 0, "3层魔法按攻击力比例(0-100)")]),
    ("安柏（火）", "兔兔伯爵", "BaronBunnyUltimate",
     "指定格模式(Tile)，在空格投放傀儡（puppetPieceId 指向傀儡资产；嘲讽/引爆由傀儡 PuppetPassive 驱动）",
     [("puppetPieceId", "int", "傀儡棋子id", 14, "傀儡在 PieceRegistry 中的 id")]),
    ("莫娜（水）", "星异", "OmenUltimate",
     "指定格模式(Tile)，目标格周围 radius 格内敌方施加易伤（受伤 +percent%），持续 duration 回合",
     [("radius", "int", "易伤半径", 2, "以目标格为中心的半径(格)"),
      ("duration", "int", "持续回合", 4, "易伤持续回合数"),
      ("percent", "int", "增伤百分比", 50, "受到的伤害提高百分比")]),
    ("爱可菲（冰）", "极寒飨宴", "FrostFeastUltimate",
     "自身增益(Self)三段：周围1格敌方物理 + 友方治疗(自身翻倍) + 激活风暴 stormTurns 回合",
     [("radius", "int", "即时段半径", 1, "以自身为中心的即时段半径(格)"),
      ("attackPercent", "int", "攻击力百分比", 100, "即时段物理攻击力比例(0-100)"),
      ("healAmount", "int", "回血量", 10, "友方回血量(自身翻倍)"),
      ("stormTurns", "int", "风暴持续回合", 4, "凛冬风暴持续回合数")]),
    ("雷电将军（雷）", "无想一刀", "MusouStrikeUltimate",
     "指定格模式(Tile)，横向3格物理伤害 + 激活协同状态 synergyTurns 回合（普攻回能 + 协同攻击）",
     [("attackPercent", "int", "攻击力百分比", 100, "横向段物理攻击力比例(0-100)"),
      ("synergyTurns", "int", "协同持续回合", 6, "协同状态持续回合数")]),
]

# ---------- 移动配置 ----------
# (棋子, providerClassName, providerJsonParams, baseRange, 说明)
MOVES = [
    ("杜林（火）", "FreeMoveProvider", "", 3, "通用移动"),
    ("菲林斯（雷）", "FreeMoveProvider", "", 4, "通用移动"),
    ("克洛琳德（雷）", "StraightPassProvider", "", 4, "直线穿越（穿棋子）"),
    ("哥伦比亚（水）", "ParityMoveProvider", "", 4, "跳格移动（偶数距离）"),
    ("甘雨（冰）", "FreeMoveProvider", "", 3, "通用移动"),
    ("安柏（火）", "FreeMoveProvider", "", 3, "通用移动"),
    ("兔兔伯爵（傀儡）", "FreeMoveProvider", "", 0, "不可移动（召唤物）"),
    ("莫娜（水）", "FlyMoveProvider", "", 3, "飞行移动（不被阻拦）"),
    ("爱可菲（冰）", "FreeMoveProvider", "", 3, "通用移动"),
    ("雷电将军（雷）", "FreeMoveProvider", "", 3, "通用移动"),
]

# ---------- 攻击配置 ----------
# (棋子, providerClassName, providerJsonParams, baseRange, 说明)
ATTACKS = [
    ("杜林（火）", "LineAttackProvider", "", 3, "直线穿透普攻"),
    ("菲林斯（雷）", "CircleAttackProvider", "", 2, "通用普攻单目标"),
    ("克洛琳德（雷）", "CircleAttackProvider", "", 1, "通用普攻单目标"),
    ("哥伦比亚（水）", "AreaAttackProvider",
     '{"splashRadius":1,"filterByAttackRange":true,"splashIsElemental":false,"splashGoldPercent":0,"baseDamage":0,"attackPercent":0.5}',
     2, "目标格+周围1格溅射"),
    ("甘雨（冰）", "CircleAttackProvider", "", 4, "远程狙击（蓄力选目标用）"),
    ("安柏（火）", "CircleAttackProvider", "", 4, "远程普攻"),
    ("兔兔伯爵（傀儡）", "CircleAttackProvider", "", 1, "不可攻击（占位）"),
    ("莫娜（水）", "CircleAttackProvider", "", 4, "远程普攻"),
    ("爱可菲（冰）", "CircleAttackProvider", "", 3, "通用普攻单目标"),
    ("雷电将军（雷）", "CircleAttackProvider", "", 3, "通用普攻单目标"),
]

# ---------- 效果速查 ----------
# (类名, 中文名, 类型, 参数, 说明)
LOOKUP = [
    ("FlameNovaPassive", "炎星溅射", "被动", "baseDamage, attackPercent, splashRadius, goldPercent, elemental", "杜林：普攻后自身周围2格溅射物理(火)"),
    ("SparkEmpowerPassive", "雷蚀强化", "被动", "duration", "菲林斯：普攻后强化6回合"),
    ("VoltPathPassive", "雷径惩戒", "被动", "baseDamage, attackPercent, elemental", "克洛琳德：移动路径飞越敌人物理伤害(雷)"),
    ("RevivePassive", "生命之息", "被动", "reviveHp", "哥伦比亚：己方死亡原地复活(一局一次)"),
    ("DoubleStrikePassive", "快速射击", "被动", "percent", "安柏：普攻两次物理伤害"),
    ("PuppetPassive", "傀儡核心", "被动", "hitsToDetonate, maxTurns, tauntRange, explosionRadius, baseDamage, attackPercent", "兔兔伯爵：受击/回合计数引爆+嘲讽光环"),
    ("DomainPassive", "水域领域", "被动", "radius, baseDamage, attackPercent", "莫娜：普攻落领域，每回合领域内敌方掉血(水)"),
    ("FrostBreakPassive", "冰锋蚀甲", "被动", "amount, duration", "爱可菲：攻击降防(不可叠加)"),
    ("BlizzardPassive", "凛冬风暴", "被动", "radius, baseDamage, attackPercent", "爱可菲：大招持续段，每回合最近/最远各1魔法(冰)"),
    ("StormEyePassive", "雷罚恶曜之眼", "被动", "energyGain", "雷电将军：队友开大回能(不含自己)"),
    ("CoordinatedStrikePassive", "协同状态", "被动", "baseDamage, attackPercent, energyPerAttack", "雷电将军：普攻回能+己方攻击触发协同"),
    ("FlameLanceUltimate", "贯日炎枪", "大招", "length, bonusDamage", "杜林：指定方向直线6格穿透(火)"),
    ("ThunderSweepUltimate", "雷殛横光", "大招", "normalBonusDamage, empoweredBaseDamage, empoweredAttackPercent, empoweredEnergyCost, empoweredElemental", "菲林斯：横向3格/强化态周围1格(雷)"),
    ("ThunderDashUltimate", "雷驰突进", "大招", "healAmount", "克洛琳德：直线冲刺+回血10"),
    ("TideSurgeUltimate", "潮涌战意", "大招", "radius, bonusAttack, duration", "哥伦比亚：周围5格己方+攻4回合"),
    ("FrostArrowUltimate", "霜华贯矢", "大招", "physicalBonusDamage, magicBaseDamage, magicAttackPercent", "甘雨：蓄力三档(1单点/2周围1/3周围2)"),
    ("BaronBunnyUltimate", "兔兔伯爵", "大招", "puppetPieceId", "安柏：投放傀儡(嘲讽+引爆)"),
    ("OmenUltimate", "星异", "大招", "radius, duration, percent", "莫娜：目标周围易伤(受伤+50%)"),
    ("FrostFeastUltimate", "极寒飨宴", "大招", "radius, attackPercent, healAmount, stormTurns", "爱可菲：伤害+治疗+激活风暴"),
    ("MusouStrikeUltimate", "无想一刀", "大招", "attackPercent, synergyTurns", "雷电将军：横向3格+激活协同状态"),
]


def build_json_formula(cells, fields):
    parts = ['="{"']
    for i, (cell, field) in enumerate(zip(cells, fields)):
        if i > 0:
            parts.append('","')
        parts.append('&CHAR(34)&"' + field + '"&CHAR(34)&":"&')
        parts.append(cell + '&')
    parts.append('"}"')
    return ''.join(parts)


def build_config_sheet(wb, title, subtitle, effects, first):
    """生成「配置生成器」sheet：多个效果块垂直堆叠，每块参数表 + className + jsonParams 公式。"""
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    for col, w in {"A": 22, "B": 12, "C": 18, "D": 12, "E": 48}.items():
        ws.column_dimensions[col].width = w

    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws.merge_cells("A1:E1")
    ws["A1"].fill = FILL_TITLE
    ws["A2"] = "使用说明：在【填写值】列（黄色底）填数值，下方【jsonParams】自动生成配置 JSON；把 className 和 jsonParams 复制到 Unity 对应配置即可。"
    ws["A2"].font = F_SUB
    ws.merge_cells("A2:E2")
    ws["A3"] = subtitle
    ws["A3"].font = F_SUB
    ws.merge_cells("A3:E3")

    row = 5
    for piece, name, cls, desc, params in effects:
        c = ws.cell(row, 1, f"◆ {piece} · {name}  {cls}")
        c.font = F_BLOCK
        for cc in range(1, 6):
            ws.cell(row, cc).fill = FILL_BLOCK
        ws.cell(row, 5, desc).font = Font(name="微软雅黑", size=9, color="FFFFFF")
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=5)

        heads = ["参数", "类型", "填写值", "默认值", "说明"]
        for i, h in enumerate(heads, 1):
            cell = ws.cell(row + 1, i, h)
            cell.font = F_HEAD
            cell.fill = FILL_HEAD
            cell.alignment = CENTER
            cell.border = BORDER

        r = row + 2
        first_param_row = r
        for (field, typ, label, default, note) in params:
            a = ws.cell(r, 1, field)
            a.font = F_FIELD
            a.border = BORDER
            ws.cell(r, 2, typ).alignment = CENTER
            ws.cell(r, 2).border = BORDER
            inp = ws.cell(r, 3, default)
            inp.fill = FILL_INPUT
            inp.alignment = CENTER
            inp.border = BORDER
            ws.cell(r, 4, default).fill = FILL_DEF
            ws.cell(r, 4).alignment = CENTER
            ws.cell(r, 4).border = BORDER
            ws.cell(r, 5, note).border = BORDER
            ws.cell(r, 5).alignment = LEFT
            r += 1

        # className 行
        ws.cell(r, 1, "className").font = Font(name="微软雅黑", size=10, bold=True)
        cls_cell = ws.cell(r, 3, cls)
        cls_cell.font = F_GEN
        cls_cell.fill = FILL_GEN
        cls_cell.alignment = CENTER
        cls_cell.border = BORDER
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        for cc in (3, 4, 5):
            ws.cell(r, cc).fill = FILL_GEN

        # jsonParams 行（公式）
        r2 = r + 1
        fields = [pp[0] for pp in params]
        cells = [f"C{first_param_row + i}" for i in range(len(fields))]
        formula = build_json_formula(cells, fields)
        ws.cell(r2, 1, "jsonParams").font = Font(name="微软雅黑", size=10, bold=True)
        jp = ws.cell(r2, 3, formula)
        jp.font = F_GEN
        jp.fill = FILL_GEN
        jp.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        jp.border = BORDER
        ws.merge_cells(start_row=r2, start_column=3, end_row=r2, end_column=5)
        for cc in (3, 4, 5):
            ws.cell(r2, cc).fill = FILL_GEN
        ws.row_dimensions[r2].height = 30

        row = r2 + 2  # 下一块


def build_range_sheet(wb, title, subtitle, configs, first):
    """生成「移动/攻击配置」sheet：内联字段表格（providerClassName + jsonParams + baseRange）。"""
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    for col, w in {"A": 16, "B": 24, "C": 46, "D": 12, "E": 26}.items():
        ws.column_dimensions[col].width = w
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws.merge_cells("A1:E1")
    ws["A1"].fill = FILL_TITLE
    ws["A2"] = subtitle
    ws["A2"].font = F_SUB
    ws.merge_cells("A2:E2")

    heads = ["棋子", "providerClassName", "providerJsonParams", "baseRange", "说明"]
    for i, h in enumerate(heads, 1):
        cell = ws.cell(3, i, h)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = CENTER
        cell.border = BORDER
    for ri, (piece, cls, jsonp, base, note) in enumerate(configs, start=4):
        vals = [piece, cls, jsonp or "（空）", base, note]
        for ci, val in enumerate(vals, 1):
            cell = ws.cell(ri, ci, val)
            cell.font = F_BODY
            cell.border = BORDER
            cell.alignment = CENTER if ci in (1, 2, 4) else LEFT
    ws.freeze_panes = "A4"


def build_lookup_sheet(wb, title, lookup):
    ws = wb.create_sheet(title)
    for col, width in {"A": 26, "B": 18, "C": 10, "D": 46, "E": 46}.items():
        ws.column_dimensions[col].width = width
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws.merge_cells("A1:E1")
    ws["A1"].fill = FILL_TITLE
    heads = ["类名", "中文名", "类型", "参数", "说明"]
    for i, h in enumerate(heads, 1):
        cell = ws.cell(2, i, h)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = CENTER
        cell.border = BORDER
    for ri, rec in enumerate(lookup, start=3):
        for ci, val in enumerate(rec, 1):
            cell = ws.cell(ri, ci, val)
            cell.font = F_BODY
            cell.border = BORDER
            cell.alignment = LEFT if ci in (1, 2, 4, 5) else CENTER
    ws.freeze_panes = "A3"


# ---------- 配置清单 ----------
# 每个棋子一块：(块标题, [(字段, 值, 说明)])
ASSET_BLOCKS = [
    ("杜林（火）· id=10", [
        ("id", "10", "注册表查找键（唯一）"),
        ("displayName", "杜林", "棋子显示名"),
        ("innateElement", "Fire", "火元素，攻击触发反应"),
        ("moveConfig.providerClassName", "FreeMoveProvider", "通用移动"),
        ("moveConfig.baseRange", "3", "移动范围"),
        ("attackConfig.providerClassName", "LineAttackProvider", "直线穿透普攻"),
        ("attackConfig.baseRange", "3", "普攻直线长度"),
        ("builtInPassives[0].className", "FlameNovaPassive", "炎星溅射被动"),
        ("builtInPassives[0].jsonParams", '{"baseDamage":0,"attackPercent":50,"splashRadius":2,"goldPercent":0,"elemental":true}', "被动参数"),
        ("ultimateConfig.targetMode", "Tile", "指定格瞄准"),
        ("ultimateConfig.tileTargetRange", "6", "瞄准半径(与 length 一致)"),
        ("ultimateConfig.effectClassName", "FlameLanceUltimate", "贯日炎枪"),
        ("ultimateConfig.effectJsonParams", '{"length":6,"bonusDamage":0}', "大招参数"),
    ]),
    ("菲林斯（雷）· id=1", [
        ("id", "1", "注册表查找键（唯一）"),
        ("displayName", "菲林斯", "棋子显示名"),
        ("innateElement", "Thunder", "雷元素"),
        ("moveConfig.providerClassName", "FreeMoveProvider", "通用移动"),
        ("moveConfig.baseRange", "4", "移动范围"),
        ("attackConfig.providerClassName", "CircleAttackProvider", "通用普攻单目标"),
        ("attackConfig.baseRange", "2", "普攻范围"),
        ("builtInPassives[0].className", "SparkEmpowerPassive", "雷蚀强化被动"),
        ("builtInPassives[0].jsonParams", '{"duration":6}', "强化持续6回合"),
        ("ultimateConfig.targetMode", "Tile", "指定格瞄准"),
        ("ultimateConfig.tileTargetRange", "3", "瞄准半径"),
        ("ultimateConfig.tileTargetShape", "Circle", "圆形高亮"),
        ("ultimateConfig.effectClassName", "ThunderSweepUltimate", "雷殛横光"),
        ("ultimateConfig.effectJsonParams", '{"normalBonusDamage":0,"empoweredBaseDamage":0,"empoweredAttackPercent":0,"empoweredEnergyCost":50,"empoweredElemental":true}', "大招参数"),
    ]),
    ("克洛琳德（雷）· id=11", [
        ("id", "11", "注册表查找键（唯一）"),
        ("displayName", "克洛琳德", "棋子显示名"),
        ("innateElement", "Thunder", "雷元素"),
        ("moveConfig.providerClassName", "StraightPassProvider", "直线穿越（穿棋子）"),
        ("moveConfig.baseRange", "4", "移动范围"),
        ("attackConfig.providerClassName", "CircleAttackProvider", "通用普攻单目标"),
        ("attackConfig.baseRange", "1", "普攻范围"),
        ("builtInPassives[0].className", "VoltPathPassive", "雷径惩戒被动"),
        ("builtInPassives[0].jsonParams", '{"baseDamage":0,"attackPercent":0,"elemental":true}', "被动参数"),
        ("ultimateConfig.targetMode", "Tile", "指定格瞄准"),
        ("ultimateConfig.tileTargetRange", "6", "冲刺射程"),
        ("ultimateConfig.tileTargetShape", "Ray", "直线冲刺(关键)"),
        ("ultimateConfig.effectClassName", "ThunderDashUltimate", "雷驰突进"),
        ("ultimateConfig.effectJsonParams", '{"healAmount":10}', "回血10"),
    ]),
    ("哥伦比亚（水）· id=12", [
        ("id", "12", "注册表查找键（唯一）"),
        ("displayName", "哥伦比亚", "棋子显示名"),
        ("innateElement", "Water", "水元素"),
        ("moveConfig.providerClassName", "ParityMoveProvider", "跳格移动（偶数距离）"),
        ("moveConfig.baseRange", "4", "移动范围"),
        ("attackConfig.providerClassName", "AreaAttackProvider", "溅射普攻"),
        ("attackConfig.providerJsonParams", '{"splashRadius":1,"filterByAttackRange":true,"splashIsElemental":false,"splashGoldPercent":0,"baseDamage":0,"attackPercent":0.5}', "溅射参数"),
        ("attackConfig.baseRange", "2", "普攻范围"),
        ("builtInPassives[0].className", "RevivePassive", "生命之息被动"),
        ("builtInPassives[0].jsonParams", '{"reviveHp":10}', "复活10血"),
        ("ultimateConfig.targetMode", "Self", "自身增益(无需指定目标)"),
        ("ultimateConfig.effectClassName", "TideSurgeUltimate", "潮涌战意"),
        ("ultimateConfig.effectJsonParams", '{"radius":5,"bonusAttack":0,"duration":4}', "大招参数"),
    ]),
    ("甘雨（冰）· id=18（asset 未建，待用户自建）", [
        ("id", "18", "建议 18（18 号，用户自建 asset 时定）"),
        ("displayName", "甘雨", "棋子显示名"),
        ("innateElement", "Ice", "冰元素"),
        ("moveConfig.providerClassName", "FreeMoveProvider", "通用移动"),
        ("moveConfig.baseRange", "3", "移动范围"),
        ("attackConfig.providerClassName", "CircleAttackProvider", "远程狙击（蓄力选目标用）"),
        ("attackConfig.baseRange", "4", "普攻选目标范围"),
        ("usesChargeSystem", "true", "关键开关：无能量条+普攻蓄力+大招按层数"),
        ("maxChargeStacks", "3", "蓄力层数上限"),
        ("builtInPassives", "（空）", "蓄力是系统开关 usesChargeSystem，不是被动"),
        ("ultimateConfig.targetMode", "Tile", "指定格瞄准"),
        ("ultimateConfig.tileTargetRange", "6", "狙击射程"),
        ("ultimateConfig.effectClassName", "FrostArrowUltimate", "霜华贯矢"),
        ("ultimateConfig.effectJsonParams", '{"physicalBonusDamage":0,"magicBaseDamage":0,"magicAttackPercent":0}', "大招参数"),
    ]),
    ("安柏（火）· id=13", [
        ("id", "13", "注册表查找键（唯一）"),
        ("displayName", "安柏", "棋子显示名"),
        ("innateElement", "Fire", "火元素"),
        ("moveConfig.providerClassName", "FreeMoveProvider", "通用移动"),
        ("moveConfig.baseRange", "3", "移动范围"),
        ("attackConfig.providerClassName", "CircleAttackProvider", "远程普攻"),
        ("attackConfig.baseRange", "4", "普攻范围"),
        ("builtInPassives[0].className", "DoubleStrikePassive", "快速射击被动"),
        ("builtInPassives[0].jsonParams", '{"percent":100}', "二段100%攻击力"),
        ("ultimateConfig.targetMode", "Tile", "指定格瞄准"),
        ("ultimateConfig.tileTargetRange", "5", "投放半径"),
        ("ultimateConfig.effectClassName", "BaronBunnyUltimate", "兔兔伯爵"),
        ("ultimateConfig.effectJsonParams", '{"puppetPieceId":14}', "指向傀儡资产 id"),
    ]),
    ("兔兔伯爵（傀儡）· id=14", [
        ("id", "14", "注册表查找键（唯一）"),
        ("displayName", "兔兔伯爵", "傀儡显示名"),
        ("isSummon", "true", "关键：临时召唤物，不计胜负/不可选中/不进选人界面"),
        ("maxHP", "9999", "高 HP 防伤害致死（受击计数不看伤害值）"),
        ("innateElement", "None", "傀儡无元素"),
        ("moveConfig.baseRange", "0", "不可移动"),
        ("attackConfig.baseRange", "1", "不可攻击（占位）"),
        ("builtInPassives[0].className", "PuppetPassive", "傀儡核心被动"),
        ("builtInPassives[0].jsonParams", '{"hitsToDetonate":3,"maxTurns":6,"tauntRange":2,"explosionRadius":1,"baseDamage":0,"attackPercent":100}', "傀儡参数"),
        ("ultimateConfig", "（空）", "无大招、无能量条"),
    ]),
    ("莫娜（水）· id=15", [
        ("id", "15", "注册表查找键（唯一）"),
        ("displayName", "莫娜", "棋子显示名"),
        ("innateElement", "Water", "水元素"),
        ("moveConfig.providerClassName", "FlyMoveProvider", "飞行移动（不被阻拦）"),
        ("moveConfig.baseRange", "3", "移动范围"),
        ("attackConfig.providerClassName", "CircleAttackProvider", "远程普攻"),
        ("attackConfig.baseRange", "4", "普攻范围"),
        ("builtInPassives[0].className", "DomainPassive", "水域领域被动"),
        ("builtInPassives[0].jsonParams", '{"radius":2,"baseDamage":0,"attackPercent":50}', "领域参数"),
        ("ultimateConfig.targetMode", "Tile", "指定格瞄准"),
        ("ultimateConfig.tileTargetRange", "5", "瞄准半径"),
        ("ultimateConfig.effectClassName", "OmenUltimate", "星异"),
        ("ultimateConfig.effectJsonParams", '{"radius":2,"duration":4,"percent":50}', "易伤参数"),
    ]),
    ("爱可菲（冰）· id=16", [
        ("id", "16", "注册表查找键（唯一）"),
        ("displayName", "爱可菲", "棋子显示名"),
        ("innateElement", "Ice", "冰元素"),
        ("moveConfig.providerClassName", "FreeMoveProvider", "通用移动"),
        ("moveConfig.baseRange", "3", "移动范围"),
        ("attackConfig.providerClassName", "CircleAttackProvider", "通用普攻单目标"),
        ("attackConfig.baseRange", "3", "普攻范围"),
        ("builtInPassives[0].className", "FrostBreakPassive", "冰锋蚀甲被动"),
        ("builtInPassives[0].jsonParams", '{"amount":4,"duration":4}', "减防4点持续4回合"),
        ("builtInPassives[1].className", "BlizzardPassive", "凛冬风暴被动"),
        ("builtInPassives[1].jsonParams", '{"radius":3,"baseDamage":0,"attackPercent":30}', "风暴参数"),
        ("ultimateConfig.targetMode", "Self", "自身增益(无需指定目标)"),
        ("ultimateConfig.effectClassName", "FrostFeastUltimate", "极寒飨宴"),
        ("ultimateConfig.effectJsonParams", '{"radius":1,"attackPercent":100,"healAmount":10,"stormTurns":4}', "大招参数"),
    ]),
    ("雷电将军（雷）· id=17", [
        ("id", "17", "注册表查找键（唯一）"),
        ("displayName", "雷电将军", "棋子显示名"),
        ("innateElement", "Thunder", "雷元素"),
        ("moveConfig.providerClassName", "FreeMoveProvider", "通用移动"),
        ("moveConfig.baseRange", "3", "移动范围"),
        ("attackConfig.providerClassName", "CircleAttackProvider", "通用普攻单目标"),
        ("attackConfig.baseRange", "3", "普攻范围"),
        ("builtInPassives[0].className", "StormEyePassive", "雷罚恶曜之眼被动"),
        ("builtInPassives[0].jsonParams", '{"energyGain":10}', "队友开大回10能量"),
        ("builtInPassives[1].className", "CoordinatedStrikePassive", "协同状态被动"),
        ("builtInPassives[1].jsonParams", '{"baseDamage":0,"attackPercent":50,"energyPerAttack":5}', "协同参数"),
        ("ultimateConfig.targetMode", "Tile", "指定格瞄准"),
        ("ultimateConfig.tileTargetRange", "3", "瞄准半径"),
        ("ultimateConfig.effectClassName", "MusouStrikeUltimate", "无想一刀"),
        ("ultimateConfig.effectJsonParams", '{"attackPercent":100,"synergyTurns":6}', "大招参数"),
    ]),
]


def build_asset_sheet(wb, title, blocks):
    ws = wb.create_sheet(title)
    for col, width in {"A": 40, "B": 52, "C": 40}.items():
        ws.column_dimensions[col].width = width
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws.merge_cells("A1:C1")
    ws["A1"].fill = FILL_TITLE
    ws["A2"] = "九个棋子 + 兔兔伯爵傀儡的 PieceData.asset 在 Unity Inspector 中的字段填写清单（内联配置，无 SO 资产）"
    ws["A2"].font = F_SUB
    ws.merge_cells("A2:C2")

    row = 3
    for piece_name, rows in blocks:
        c = ws.cell(row, 1, f"◆ {piece_name}")
        c.font = F_BLOCK
        for cc in range(1, 4):
            ws.cell(row, cc).fill = FILL_BLOCK
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        row += 1
        for (field, val, note) in rows:
            ws.cell(row, 1, field).font = F_FIELD
            ws.cell(row, 1).border = BORDER
            ws.cell(row, 2, val).font = F_BODY
            ws.cell(row, 2).border = BORDER
            ws.cell(row, 2).alignment = LEFT
            ws.cell(row, 3, note).font = F_SUB
            ws.cell(row, 3).border = BORDER
            ws.cell(row, 3).alignment = LEFT
            row += 1
        row += 1  # 块间空一行
    ws.freeze_panes = "A3"


# ============================================================
#  生成 6 张表
# ============================================================
wb = openpyxl.Workbook()
build_config_sheet(wb, "棋子被动配置生成器", "九个棋子的内建被动配置（含傀儡 PuppetPassive；带元素参数，默认附着先天元素触发反应）", PASSIVES, first=True)
build_config_sheet(wb, "棋子大招配置生成器", "九棋子的大招配置", ULTIMATES, first=False)
build_range_sheet(wb, "移动配置生成器", "移动范围内联配置（直接写入 PieceData.moveConfig）", MOVES, first=False)
build_range_sheet(wb, "攻击配置生成器", "攻击范围内联配置（直接写入 PieceData.attackConfig）", ATTACKS, first=False)
build_lookup_sheet(wb, "棋子效果速查", LOOKUP)
build_asset_sheet(wb, "棋子配置清单", ASSET_BLOCKS)

wb.save(OUT)
print("SAVED:", OUT)
